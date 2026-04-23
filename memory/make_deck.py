"""Generate QPGEN product presentation (.pptx).

Swiss / brutalist light theme matching the web app:
- Off-white background
- Klein blue #002FA7 accent
- Yellow #FFC300 highlight
- Black hard-edged borders and accents
- IBM Plex Sans / Calibri fallbacks
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Colors
INK = RGBColor(0x0A, 0x0A, 0x0A)
INK_SOFT = RGBColor(0x52, 0x52, 0x52)
BG = RGBColor(0xFA, 0xFA, 0xFA)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x00, 0x2F, 0xA7)
YELLOW = RGBColor(0xFF, 0xC3, 0x00)
RED = RGBColor(0xE6, 0x39, 0x46)
GREEN = RGBColor(0x10, 0xB9, 0x81)

FONT_HEADING = "Calibri"  # Cabinet Grotesk is not system-wide; fallback
FONT_BODY = "Calibri"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_filled_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=INK,
    font=FONT_BODY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    tracking=None,
):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if tracking:
        run.font._element.set("spc", str(tracking))
    return tb


def add_overline(slide, left, top, text, width=Inches(4.5)):
    add_text(
        slide,
        left,
        top,
        width,
        Inches(0.25),
        text,
        size=10,
        bold=True,
        color=INK_SOFT,
        font=FONT_MONO,
    )


def add_heading(slide, left, top, text, size=44, width=Inches(11)):
    add_text(
        slide,
        left,
        top,
        width,
        Inches(1.2),
        text,
        size=size,
        bold=True,
        color=INK,
        font=FONT_HEADING,
    )


def add_body_bullets(slide, left, top, width, height, items, size=14):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = f"→  {item}"
        run.font.name = FONT_BODY
        run.font.size = Pt(size)
        run.font.color.rgb = INK


def background(slide):
    add_filled_rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)
    # top black bar
    add_filled_rect(slide, 0, 0, SLIDE_W, Inches(0.06), INK)
    # left thick stripe
    add_filled_rect(slide, 0, 0, Inches(0.08), SLIDE_H, INK)
    # footer
    add_text(
        slide,
        Inches(0.4),
        Inches(7.1),
        Inches(8),
        Inches(0.3),
        "QPGEN — AI Question Paper Studio",
        size=9,
        color=INK_SOFT,
        font=FONT_MONO,
    )
    add_text(
        slide,
        Inches(11.4),
        Inches(7.1),
        Inches(1.5),
        Inches(0.3),
        "// CONFIDENTIAL",
        size=9,
        color=INK_SOFT,
        font=FONT_MONO,
        align=PP_ALIGN.RIGHT,
    )


def page_number(slide, n):
    add_text(
        slide,
        Inches(12.6),
        Inches(0.3),
        Inches(0.5),
        Inches(0.3),
        f"{n:02d}",
        size=11,
        bold=True,
        color=INK,
        font=FONT_MONO,
    )


def make_card(slide, left, top, width, height, title, body_lines, accent=BLUE):
    # outer black border rect
    shape = add_filled_rect(slide, left, top, width, height, SURFACE, line=INK)
    # left accent strip
    add_filled_rect(slide, left, top, Inches(0.14), height, accent)
    # title
    add_text(
        slide,
        left + Inches(0.35),
        top + Inches(0.25),
        width - Inches(0.5),
        Inches(0.35),
        title.upper(),
        size=12,
        bold=True,
        color=INK,
        font=FONT_MONO,
        tracking=200,
    )
    # body
    add_body_bullets(
        slide,
        left + Inches(0.35),
        top + Inches(0.7),
        width - Inches(0.5),
        height - Inches(1),
        body_lines,
        size=12,
    )


# ---------- Slides ----------

def slide_cover(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    # Big decorative black logo block
    add_filled_rect(s, Inches(0.6), Inches(0.55), Inches(0.9), Inches(0.9), INK)
    add_text(
        s,
        Inches(0.6),
        Inches(0.7),
        Inches(0.9),
        Inches(0.6),
        "Q",
        size=44,
        bold=True,
        color=YELLOW,
        font=FONT_HEADING,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        s,
        Inches(1.7),
        Inches(0.6),
        Inches(8),
        Inches(0.4),
        "QPGEN",
        size=16,
        bold=True,
        color=INK,
        font=FONT_HEADING,
    )
    add_text(
        s,
        Inches(1.7),
        Inches(1),
        Inches(8),
        Inches(0.3),
        "// QUESTION PAPER STUDIO",
        size=10,
        bold=True,
        color=INK_SOFT,
        font=FONT_MONO,
    )

    add_overline(s, Inches(0.6), Inches(2.3), "// VOLUME 01 · PRODUCT OVERVIEW")

    # Big title
    tb = s.shapes.add_textbox(Inches(0.6), Inches(2.65), Inches(12), Inches(3))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    r = p1.add_run()
    r.text = "Design exams"
    r.font.name = FONT_HEADING
    r.font.size = Pt(72)
    r.font.bold = True
    r.font.color.rgb = INK
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "algorithmically."
    r.font.name = FONT_HEADING
    r.font.size = Pt(72)
    r.font.bold = True
    r.font.color.rgb = BLUE

    add_text(
        s,
        Inches(0.6),
        Inches(5.45),
        Inches(9),
        Inches(0.5),
        "An AI-powered question paper generator for schools and colleges.",
        size=18,
        color=INK_SOFT,
    )
    add_text(
        s,
        Inches(0.6),
        Inches(5.85),
        Inches(9),
        Inches(0.4),
        "Upload textbooks → extract topics → generate original papers, answer keys and diagrams.",
        size=14,
        color=INK_SOFT,
    )

    # Right info block
    add_filled_rect(s, Inches(10), Inches(2.5), Inches(2.75), Inches(3.5), INK)
    add_text(
        s,
        Inches(10.2),
        Inches(2.7),
        Inches(2.5),
        Inches(0.3),
        "// STACK",
        size=9,
        bold=True,
        color=YELLOW,
        font=FONT_MONO,
    )
    add_body_bullets(
        s,
        Inches(10.2),
        Inches(3.0),
        Inches(2.5),
        Inches(3),
        [
            "FastAPI",
            "React + Tailwind",
            "MongoDB",
            "OpenAI GPT-5.2",
            "Claude Sonnet 4.5",
            "Nano Banana (diagrams)",
            "Emergent Object Store",
        ],
        size=11,
    )


def slide_section(prs, n, number, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), f"// SECTION {number:02d}")
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1), Inches(12), Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title.upper()
    r.font.name = FONT_HEADING
    r.font.size = Pt(88)
    r.font.bold = True
    r.font.color.rgb = INK
    p2 = tf.add_paragraph()
    p2.space_before = Pt(10)
    r = p2.add_run()
    r.text = subtitle
    r.font.name = FONT_BODY
    r.font.size = Pt(20)
    r.font.color.rgb = BLUE


def slide_problem(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), "// WHY QPGEN")
    add_heading(s, Inches(0.6), Inches(1), "The problem we solve.", size=48)

    # Two-column
    left = Inches(0.6)
    col_w = Inches(5.9)
    col_h = Inches(4.3)
    make_card(
        s,
        left,
        Inches(2.6),
        col_w,
        col_h,
        "// STATUS QUO",
        [
            "Teachers spend 4-6 hours per paper",
            "Question banks get repetitive year after year",
            "Difficulty mix and topic coverage are guesswork",
            "Answer keys are handwritten or missing",
            "No system to learn teacher preferences",
            "Diagrams drawn manually or skipped",
        ],
        accent=RED,
    )
    make_card(
        s,
        Inches(6.8),
        Inches(2.6),
        col_w,
        col_h,
        "// WITH QPGEN",
        [
            "Original paper in under 60 seconds",
            "Topics auto-extracted from any textbook PDF",
            "Exact control over difficulty & type distribution",
            "AI answer keys with step-by-step solutions",
            "Every edit feeds a learning loop",
            "Line-drawn diagrams auto-generated",
        ],
        accent=GREEN,
    )


def slide_architecture(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), "// ARCHITECTURE")
    add_heading(s, Inches(0.6), Inches(1), "How the engine works.", size=48)

    # Pipeline blocks
    x = Inches(0.6)
    y = Inches(2.5)
    w = Inches(2.2)
    h = Inches(1.6)
    gap = Inches(0.15)
    steps = [
        ("01", "UPLOAD", "Textbook PDF", BLUE),
        ("02", "INDEX", "Parse + chunk + store", INK),
        ("03", "EXTRACT", "Topics via LLM", BLUE),
        ("04", "GENERATE", "Paper + diagrams + answers", YELLOW),
        ("05", "REFINE", "Teacher edits feedback loop", GREEN),
    ]
    for i, (num, title, desc, color) in enumerate(steps):
        left = x + (w + gap) * i
        add_filled_rect(s, left, y, w, h, SURFACE, line=INK)
        add_filled_rect(s, left, y, w, Inches(0.35), color)
        add_text(
            s,
            left + Inches(0.2),
            y + Inches(0.05),
            w,
            Inches(0.3),
            num,
            size=11,
            bold=True,
            color=SURFACE,
            font=FONT_MONO,
        )
        add_text(
            s,
            left + Inches(0.2),
            y + Inches(0.5),
            w,
            Inches(0.4),
            title,
            size=16,
            bold=True,
            color=INK,
            font=FONT_HEADING,
        )
        add_text(
            s,
            left + Inches(0.2),
            y + Inches(1),
            w - Inches(0.4),
            Inches(0.6),
            desc,
            size=11,
            color=INK_SOFT,
        )

    # Bottom: provider fallback
    y2 = Inches(4.8)
    add_overline(s, Inches(0.6), y2, "// LLM ADAPTER — AUTOMATIC FAILOVER")
    add_filled_rect(s, Inches(0.6), y2 + Inches(0.35), Inches(12.1), Inches(1.5), INK)
    add_text(
        s,
        Inches(0.9),
        y2 + Inches(0.55),
        Inches(12),
        Inches(0.4),
        "OpenAI GPT-5.2  →  Claude Sonnet 4.5  →  (future: open-source models)",
        size=18,
        bold=True,
        color=YELLOW,
        font=FONT_MONO,
    )
    add_text(
        s,
        Inches(0.9),
        y2 + Inches(1.1),
        Inches(12),
        Inches(0.4),
        "Paper generation automatically falls back to the next provider on budget/rate/outage errors.",
        size=12,
        color=SURFACE,
    )


def slide_personas(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), "// USERS")
    add_heading(s, Inches(0.6), Inches(1), "Who uses QPGEN.", size=48)

    cards = [
        (
            "TEACHER",
            [
                "Uploads textbooks",
                "Selects topics & difficulty",
                "Tunes question-type mix",
                "Edits questions & answers",
                "Downloads printable PDFs",
                "Saves favourites to bank",
            ],
            BLUE,
        ),
        (
            "ADMIN",
            [
                "Manages subjects & classes",
                "Bulk-uploads textbooks",
                "Uploads prior question papers",
                "Sees all institution papers",
                "Triggers bulk solution generation",
                "Controls distribution standards",
            ],
            INK,
        ),
        (
            "STUDENT (MVP-lite)",
            [
                "Selects topics to revise",
                "Gets a timed practice paper",
                "Views solution after attempt",
                "Tracks weak topics",
                "Offline PDF export",
                "Roadmapped for next phase",
            ],
            INK_SOFT,
        ),
    ]
    x = Inches(0.6)
    y = Inches(2.4)
    w = Inches(4.05)
    h = Inches(4.4)
    gap = Inches(0.15)
    for i, (title, items, color) in enumerate(cards):
        make_card(s, x + (w + gap) * i, y, w, h, title, items, accent=color)


def slide_page(prs, n, overline, title, sub, how_it_works, utility, route):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), overline)
    add_heading(s, Inches(0.6), Inches(1), title, size=48)

    # Route badge
    route_w = Inches(4)
    add_filled_rect(s, Inches(0.6), Inches(2.05), route_w, Inches(0.45), INK)
    add_text(
        s,
        Inches(0.75),
        Inches(2.12),
        route_w,
        Inches(0.35),
        f"ROUTE  {route}",
        size=11,
        bold=True,
        color=YELLOW,
        font=FONT_MONO,
    )

    add_text(
        s,
        Inches(0.6),
        Inches(2.7),
        Inches(12.1),
        Inches(0.5),
        sub,
        size=16,
        color=INK_SOFT,
    )

    # Two cards
    make_card(
        s,
        Inches(0.6),
        Inches(3.5),
        Inches(6),
        Inches(3.3),
        "// HOW IT WORKS",
        how_it_works,
        accent=BLUE,
    )
    make_card(
        s,
        Inches(6.75),
        Inches(3.5),
        Inches(6),
        Inches(3.3),
        "// UTILITY",
        utility,
        accent=YELLOW,
    )


def slide_feature(prs, n, overline, title, sub, points):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), overline)
    add_heading(s, Inches(0.6), Inches(1), title, size=48)
    add_text(
        s,
        Inches(0.6),
        Inches(2.1),
        Inches(12),
        Inches(0.5),
        sub,
        size=16,
        color=INK_SOFT,
    )
    # Big body card
    make_card(
        s,
        Inches(0.6),
        Inches(2.9),
        Inches(12.1),
        Inches(3.9),
        "// DETAILS",
        points,
        accent=BLUE,
    )


def slide_roadmap(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), "// WHAT'S NEXT")
    add_heading(s, Inches(0.6), Inches(1), "Roadmap.", size=48)
    lanes = [
        (
            "P0 — HARDENING",
            [
                "Regenerate single question",
                "Similarity check vs. textbook",
                "Admin seed + role matrix",
            ],
            BLUE,
        ),
        (
            "P1 — GROWTH",
            [
                "School-wide shared library",
                "Student practice mode (full)",
                "Answer validation scoring",
                "Email / WhatsApp paper sharing",
            ],
            YELLOW,
        ),
        (
            "P2 — PLATFORM",
            [
                "Multi-tenant SaaS",
                "AI evaluation of scripts",
                "Difficulty auto-calibration",
                "True vector DB + RAG",
            ],
            GREEN,
        ),
    ]
    x = Inches(0.6)
    y = Inches(2.4)
    w = Inches(4.05)
    h = Inches(4.4)
    gap = Inches(0.15)
    for i, (title, items, color) in enumerate(lanes):
        make_card(s, x + (w + gap) * i, y, w, h, title, items, accent=color)


def slide_closing(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s)
    page_number(s, n)
    add_overline(s, Inches(0.6), Inches(0.6), "// THANK YOU")
    # Huge word
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12), Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Papers,"
    r.font.name = FONT_HEADING
    r.font.size = Pt(120)
    r.font.bold = True
    r.font.color.rgb = INK
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "in minutes."
    r.font.name = FONT_HEADING
    r.font.size = Pt(120)
    r.font.bold = True
    r.font.color.rgb = BLUE

    add_text(
        s,
        Inches(0.6),
        Inches(6),
        Inches(12),
        Inches(0.5),
        "Questions? Feedback? Let's make exams smarter.",
        size=18,
        color=INK_SOFT,
    )


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    n = 1
    slide_cover(prs, n); n += 1
    slide_problem(prs, n); n += 1
    slide_personas(prs, n); n += 1
    slide_architecture(prs, n); n += 1

    slide_section(prs, n, 2, "The Pages.", "A walkthrough of every screen in the product."); n += 1

    slide_page(
        prs, n,
        "// 01 · AUTH",
        "Login & Register.",
        "Email + password JWT auth with role selection. Protected routes route the user straight to the dashboard on success.",
        [
            "React form calls POST /api/auth/register or /login",
            "Backend hashes passwords with bcrypt, signs JWT (7-day)",
            "Token cached in localStorage; axios injects Bearer header",
            "Role field (teacher / admin) drives access on server",
        ],
        [
            "Zero-friction onboarding — under 10 seconds",
            "Role-based data isolation between teachers",
            "Foundation for per-institution multi-tenancy",
            "JWT means stateless API, horizontally scalable",
        ],
        "/login  ·  /register",
    ); n += 1

    slide_page(
        prs, n,
        "// 02 · DASHBOARD",
        "The control room.",
        "Landing page after login. Shows institution-wide stats, recent papers, and quick actions.",
        [
            "Parallel GET /papers, /textbooks, /qbank on load",
            "Stats cards: textbooks indexed / papers / questions saved",
            "Recent-papers list with one-click open",
            "CTA: Generate New Paper + Bulk Generate Solutions",
        ],
        [
            "Teacher knows their status at a glance",
            "Single button accelerates the core workflow",
            "Bulk solution action saves hours on legacy papers",
            "Empty states guide new users through onboarding",
        ],
        "/",
    ); n += 1

    slide_page(
        prs, n,
        "// 03 · TEXTBOOKS",
        "Knowledge ingestion.",
        "Teachers and admins upload textbook PDFs. The system parses, chunks, stores and optionally extracts topics via LLM.",
        [
            "Drag-drop PDF (up to 500 MB) with subject + class",
            "Backend saves file in Emergent Object Storage",
            "pypdf extracts text → chunked at ~2 500 chars",
            "Extract Topics button runs LLM → topic tree",
            "Status badges: indexed · topics_ready · failed",
        ],
        [
            "One upload becomes the scope + style guide forever",
            "No manual topic tagging — LLM does it",
            "Tree view lets teachers validate coverage visually",
            "Re-indexing lets you adjust syllabus alignment",
        ],
        "/textbooks",
    ); n += 1

    slide_page(
        prs, n,
        "// 04 · NEW PAPER",
        "Tune the generator.",
        "The density-heavy control-room form: metadata, difficulty, duration, marks, distribution sliders and topic multi-select.",
        [
            "Textbook dropdown + paper metadata",
            "Sliders for Information / Concept / Application %",
            "Distribution validated to sum to 100",
            "Topic chips: only select what you want covered",
            "POST /papers/generate → JSON sections with type & marks",
        ],
        [
            "Every lever a teacher cares about in one screen",
            "Instant feedback on distribution totals",
            "Feedback-loop context is sent invisibly to the LLM",
            "Diagrams are auto-requested where relevant",
        ],
        "/papers/new",
    ); n += 1

    slide_page(
        prs, n,
        "// 05 · PAPER VIEW",
        "The generated paper.",
        "Exam-paper-like preview with per-question type badges, diagrams, edit and download.",
        [
            "Sections render like a real exam paper",
            "Badges: question type, difficulty, diagram presence",
            "Per-question star (important) and bookmark actions",
            "Edit Mode: inline text, marks, type, difficulty controls",
            "Download PDF via fetch→blob with 3 s revoke delay",
        ],
        [
            "Looks printable — teachers see the final artefact",
            "Edit once, AI learns the correction for next time",
            "Saved questions feed the question bank",
            "Diagram placeholder auto-loads when ready",
        ],
        "/papers/:id",
    ); n += 1

    slide_page(
        prs, n,
        "// 06 · SOLUTION VIEW",
        "The answer key.",
        "On-demand companion to the paper: concise for info, explanatory for concept, step-by-step for application questions.",
        [
            "Generate Solution button calls /solution/generate",
            "Answer-depth differs by question type automatically",
            "Editable per-answer textarea with save",
            "Stale banner appears when the paper changes later",
            "Separate Download Solution PDF button",
        ],
        [
            "Teachers verify student answers quickly",
            "Handwritten keys replaced with structured PDF",
            "Teacher rewrites feed the solution feedback loop",
            "Clear stale flag prevents confusion on edited papers",
        ],
        "/papers/:id/solution",
    ); n += 1

    slide_page(
        prs, n,
        "// 07 · QUESTION BANK",
        "The living repository.",
        "Search, filter, save, or bulk-upload existing question papers — builds an ever-growing bank.",
        [
            "Search by keyword · filter by type & difficulty",
            "Save-from-paper promotes a single question to bank",
            "Upload Question Paper: parses any PDF → extracts qs",
            "Both teachers and admins can upload",
            "Each item tagged with source, subject, class",
        ],
        [
            "Institutional memory of every good question asked",
            "Reusable across classes and years",
            "Legacy papers become structured data instantly",
            "Accelerates new paper generation (future RAG)",
        ],
        "/qbank",
    ); n += 1

    slide_section(prs, n, 3, "The Magic.", "Under the hood features that set QPGEN apart."); n += 1

    slide_feature(
        prs, n,
        "// FEATURE A",
        "Self-learning feedback loop.",
        "Every teacher edit becomes permanent training data for their next paper.",
        [
            "PATCH /papers/:id logs each modified / added / deleted question to paper_edits",
            "On the next POST /papers/generate, recent edits are summarised and injected as 'TEACHER PREFERENCES' in the LLM prompt",
            "Same flow exists for solutions: solution_edits feeds future answer-key generation",
            "Result: the system gets measurably better per teacher, per subject, over time — without any model fine-tuning",
        ],
    ); n += 1

    slide_feature(
        prs, n,
        "// FEATURE B",
        "Auto-generated diagrams.",
        "Physics, Biology, Geography or Maths papers get clean line diagrams via Gemini Nano Banana.",
        [
            "LLM flags needs_diagram + diagram_description per question during generation",
            "Backend fires up to 5 diagram jobs in parallel (background), saves PNGs to object storage",
            "Frontend auto-polls every 4 s and swaps in images as they arrive — no refresh needed",
            "Diagrams are embedded in both the question paper PDF and the solution PDF",
        ],
    ); n += 1

    slide_feature(
        prs, n,
        "// FEATURE C",
        "Multi-provider LLM adapter.",
        "Never a single point of failure. Budget, rate, or outage on one provider automatically routes to the next.",
        [
            "Chain: OpenAI GPT-5.2 → Claude Sonnet 4.5 → (pluggable)",
            "Unified chat_complete() and parse_json_response() tolerate code fences, prose and partial JSON",
            "Adapter pattern makes adding new providers a one-line change",
            "Cost observability via Emergent Universal Key dashboard",
        ],
    ); n += 1

    slide_feature(
        prs, n,
        "// FEATURE D",
        "Bulk operations.",
        "One button, many papers. Designed for teachers with backlogs of legacy question papers.",
        [
            "POST /papers/solutions/bulk-generate?only_missing=true processes every paper without a solution",
            "Returns succeeded / failed counts and full items so teachers can retry individually",
            "Upload-question-paper workflow parses PDF → LLM extracts all questions → saves to qbank",
            "Future: admin-level dashboard to bulk-regen stale solutions nightly",
        ],
    ); n += 1

    slide_personas(prs, n); n += 1
    slide_roadmap(prs, n); n += 1
    slide_closing(prs, n); n += 1

    out = "/app/memory/QPGEN_Product_Deck.pptx"
    prs.save(out)
    print(f"Saved: {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
